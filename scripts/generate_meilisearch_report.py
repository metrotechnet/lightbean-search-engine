#!/usr/bin/env python3
"""Generate a Word report from Meilisearch chat answer results."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document


def load_json_array(path: Path) -> list[dict]:
    with path.open('r', encoding='utf-8') as fh:
        data = json.load(fh)
    if not isinstance(data, list):
        raise ValueError(f'{path} does not contain a JSON array')
    return data


def extract_question_answer(item: dict) -> tuple[str, str]:
    question = (
        item.get('question_fr')
        or item.get('question_en')
        or item.get('query')
        or item.get('question')
        or item.get('prompt')
        or ''
    )
    answer = item.get('answer') or item.get('response') or item.get('result') or ''
    return str(question), str(answer)


def build_question_answer_report(answers: list[dict], output_path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, 0)

    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Question'
    hdr[1].text = 'Réponse'

    for item in answers:
        row = table.add_row().cells
        question, answer = extract_question_answer(item)
        row[0].text = question[:1000]
        row[1].text = answer[:4000]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        if output_path.exists():
            output_path.unlink()
    except PermissionError:
        print(f'Warning: report already open in another application: {output_path}')
        return
    try:
        document.save(output_path)
    except PermissionError:
        print(f'Warning: unable to save report because it is open in another application: {output_path}')
        return
    print(f'Generated report: {output_path}')


def main() -> int:
    parser = argparse.ArgumentParser(description='Generate a Word report with only the question/answer table for a selected JSON file.')
    parser.add_argument('input_path', nargs='?', default=None, help='Path to the input JSON file containing question/answer entries')
    parser.add_argument('-o', '--output', default=None, help='Output .docx path')
    args = parser.parse_args()

    project_root = Path(__file__).resolve().parent.parent
    input_path = Path(args.input_path) if args.input_path else project_root / 'results' / 'meilisearch-chat-answers_default.json'
    input_path = input_path if input_path.is_absolute() else (project_root / input_path)

    if not input_path.exists():
        raise FileNotFoundError(f'Missing input JSON file: {input_path}')

    answers = load_json_array(input_path)
    base_name = input_path.stem
    output_path = Path(args.output) if args.output else project_root / 'results' / f'{base_name}-table.docx'
    output_path = output_path if output_path.is_absolute() else (project_root / output_path)

    build_question_answer_report(answers, output_path, f'Tableau question / réponse - {base_name}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
